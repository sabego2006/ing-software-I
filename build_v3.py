"""Build v3.docx from v2.docx with the grill decisions applied.

Decisions (grill cerrado 2026-08-24):
- RF-10: destino favorito (no ruta con hora fija).
- RF-12: registrar suspensión + notificar a usuarios que tengan la ruta como favorita/recurrente.
- RF-13: 5 métricas de uso (usuarios, búsquedas/día, top 5, comentarios pendientes, suspensiones activas).
- RNF-01: p95 < 5 s + modo offline básico (sin cálculo de tiempo de llegada, solo la ruta).
- RNF-02: 95% mensual.
- RNF-04: 4 pasos, login 1 vez/semana.
- RNF-05: solo herramientas en CI (OWASP Dependency-Check + análisis estático).
- RNF-06: bcrypt explícito.
- RNF-07: Chrome + Safari + escritorio.
- RNF-08: eliminar el RNF.
- Métricas: 5 (ISO 25010), eliminar las 3 propias.
"""

import copy
import shutil
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt

SRC = Path(
    r"C:\Users\Santiago\OneDrive - UNIVERSIDAD DE CUNDINAMARCA\Universidad\5 SEMESTRE\INGENIERIA SOFTWARE I\actividad 3 ing soft_APA v2.docx"
)
DST = Path(
    r"C:\Users\Santiago\OneDrive - UNIVERSIDAD DE CUNDINAMARCA\Universidad\5 SEMESTRE\INGENIERIA SOFTWARE I\actividad 3 ing soft_APA v3.docx"
)

shutil.copyfile(SRC, DST)
doc = docx.Document(DST)

# -------------------------------------------------------------------
# Tabla 0 (Alcance) — fila 6, 9, 10 ajustadas al grill.
# -------------------------------------------------------------------
t0 = doc.tables[0]
t0.rows[6].cells[0].text = (
    "Aviso de trancón en hora pico y destino favorito del usuario"
)
t0.rows[9].cells[0].text = (
    "CRUD de rutas, suspensión temporal y notificación a usuarios relevantes (administrador)"
)
t0.rows[10].cells[0].text = (
    "Métricas de uso: usuarios registrados, búsquedas por día, top 5 rutas, "
    "comentarios pendientes de respuesta y suspensiones activas (administrador)"
)

# -------------------------------------------------------------------
# Tabla 1 (RF) — reescribir filas 10, 12, 13 (RF-10, RF-12, RF-13).
# -------------------------------------------------------------------
t1 = doc.tables[1]


def rewrite_cell(cell, text):
    """Replace cell content keeping the first paragraph's style."""
    # Remove all paragraphs except the first.
    for p in list(cell.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    # Clean the first paragraph runs.
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.font.size = Pt(10)


# RF-10 — destino favorito
rewrite_cell(
    t1.rows[10].cells[2],
    "El usuario puede marcar un destino como favorito; la aplicación lo pre-carga "
    "al inicio de la búsqueda y muestra el aviso de trancón habitual para esa ruta "
    "a la hora actual de la consulta.",
)
# RF-12 — registrar + notificar a usuarios relevantes
rewrite_cell(
    t1.rows[12].cells[2],
    "El administrador puede suspender temporalmente una ruta (por ejemplo, en días "
    "festivos o cívicos) y, al hacerlo, el sistema envía una notificación automática "
    "a los usuarios que tengan esa ruta como favorita o destino recurrente.",
)
# RF-13 — cinco métricas
rewrite_cell(
    t1.rows[13].cells[2],
    "El administrador puede ver cinco métricas de uso: (1) total de usuarios "
    "registrados, (2) número de búsquedas por día, (3) las 5 rutas más consultadas, "
    "(4) cantidad de comentarios pendientes de respuesta y (5) suspensiones activas.",
)

# -------------------------------------------------------------------
# Tabla 2 (RNF) — reescribir o eliminar filas.
# -------------------------------------------------------------------
t2 = doc.tables[2]

# RNF-01 — p95 < 5 s + modo offline básico
rewrite_cell(
    t2.rows[1].cells[2],
    "El sistema debe responder las consultas de ruta con un p95 menor a 5 segundos "
    "medido extremo a extremo (cliente → API → render) en una red 4G típica, "
    "verificado con Spring Boot Actuator + Micrometer. La aplicación ofrece un modo "
    "offline básico que muestra la ruta sugerida y los barrios por los que pasa, "
    "sin incluir el cálculo de tiempo de llegada, el cual sí se obtiene en línea "
    "vía la API externa de geocodificación.",
)
# RNF-02 — 95% mensual
rewrite_cell(
    t2.rows[2].cells[2],
    "El sistema debe mantener al menos un 95% de disponibilidad mensual del backend, "
    "medida con UptimeRobot (latido cada 5 minutos), en horario hábil de lunes a "
    "viernes entre 7:00 y 21:00 hora Colombia. RTO < 1 hora lectiva y RPO < 24 horas, "
    "con backup diario automatizado de PostgreSQL.",
)
# RNF-03 (fila 3) — sin cambios, pero limpiamos por si acaso
# (mantener)
# RNF-04 — 4 pasos, login 1 vez/semana
rewrite_cell(
    t2.rows[4].cells[2],
    "El flujo principal de búsqueda de ruta (una vez iniciada la sesión) debe tomar "
    "máximo 4 pasos: (1) abrir la app, (2) marcar origen, (3) marcar destino, "
    "(4) ver el resultado. El inicio de sesión se solicita como máximo una vez por "
    "semana, lo que equivale a un clic extra 1 de cada 7 días. La métrica se mide "
    "con al menos 5 usuarios externos al equipo y se exige una tasa de éxito ≥ 80% "
    "en el primer intento.",
)
# RNF-05 — solo herramientas en CI
rewrite_cell(
    t2.rows[5].cells[2],
    "El sistema no debe tener vulnerabilidades altas o críticas reportadas por las "
    "herramientas automatizadas de análisis en el momento de cada entrega. Las "
    "herramientas son OWASP Dependency-Check y un analizador estático ejecutados en "
    "CI en cada Pull Request, y la métrica se reporta como '0 vulnerabilidades altas "
    "o críticas detectadas en la fecha X'.",
)
# RNF-06 — bcrypt explícito
rewrite_cell(
    t2.rows[6].cells[2],
    "Las contraseñas de los usuarios deben almacenarse con hash bcrypt (factor de "
    "costo ≥ 10), nunca en texto plano. El cumplimiento se valida con un test de "
    "integración que verifica que el hash guardado no es igual a la contraseña en "
    "plano y que el login funciona con la contraseña original.",
)
# RNF-07 — Chrome + Safari + escritorio
rewrite_cell(
    t2.rows[7].cells[2],
    "La aplicación debe funcionar correctamente en los navegadores Chrome (Android), "
    "Safari (iOS) y al menos un navegador de escritorio (Chrome o Firefox), probados "
    "en dispositivos físicos y en emulador según disponibilidad.",
)

# RNF-08 (fila 8) — eliminar
last_row = t2.rows[8]
last_row._element.getparent().remove(last_row._element)

# -------------------------------------------------------------------
# Tabla 3 (Métricas) — reemplazar 8 filas por 5 rediseñadas.
# -------------------------------------------------------------------
t3 = doc.tables[3]


def replace_row_text(row, cells_text):
    for cell, text in zip(row.cells, cells_text):
        rewrite_cell(cell, text)


# Limpiar la fila de encabezado y reescribir
replace_row_text(
    t3.rows[0],
    ["Atributo ISO/IEC 25010", "Métrica", "Instrumento de medición"],
)
# Reescribir las 5 métricas
new_metrics = [
    (
        "Eficiencia de desempeño (rendimiento)",
        "p95 < 5 s en /api/routes/search (medido extremo a extremo) "
        "+ LCP < 2.5 s en Lighthouse mobile con throttling 'Slow 4G'",
        "Spring Boot Actuator + Micrometer (latencia por endpoint, percentil 95) "
        "y Lighthouse en CI sobre la pantalla de resultado de búsqueda.",
    ),
    (
        "Fiabilidad (disponibilidad)",
        "≥ 95% de uptime mensual del backend en horario hábil (lun–vie 7:00–21:00), "
        "con RTO < 1 hora lectiva y RPO < 24 horas",
        "UptimeRobot (latido cada 5 minutos) + backup diario automatizado de "
        "PostgreSQL y restore probado al cierre de cada sprint.",
    ),
    (
        "Mantenibilidad",
        "0 violaciones de la arquitectura hexagonal detectadas en CI + "
        "≥ 70% de cobertura de tests unitarios en el paquete domain/",
        "ArchUnit ejecutándose en CI sobre los módulos del backend + reporte de "
        "cobertura JaCoCo con umbral mínimo bloqueante.",
    ),
    (
        "Usabilidad",
        "Flujo de búsqueda completado en ≤ 4 acciones del usuario (sin contar login) "
        "con tasa de éxito ≥ 80% en el primer intento",
        "Prueba con al menos 5 usuarios externos al equipo antes de cada "
        "sustentación, registrando pasos y resultado de cada intento.",
    ),
    (
        "Seguridad",
        "0 vulnerabilidades altas o críticas reportadas por herramientas "
        "automatizadas en el momento de la entrega",
        "OWASP Dependency-Check + analizador estático ejecutados en CI en cada "
        "Pull Request; resultado reportado en la sustentación.",
    ),
]

# Reemplazar las filas existentes (R1..R5) con las nuevas métricas
for i, (atributo, metrica, instrumento) in enumerate(new_metrics, start=1):
    replace_row_text(t3.rows[i], [atributo, metrica, instrumento])

# Eliminar las filas restantes (R6, R7, R8 — métricas propias)
rows_to_remove = list(t3.rows)[6:]
for row in rows_to_remove:
    row._element.getparent().remove(row._element)

# -------------------------------------------------------------------
# Insertar un párrafo al inicio (después del título) que registre la v3.
# -------------------------------------------------------------------
body = doc.element.body
# Buscar el primer párrafo (portada) y luego insertar un nuevo párrafo
# inmediatamente después de la portada, antes de "DEFINICIÓN DEL PROYECTO..."
target_text = "DEFINICIÓN DEL PROYECTO INTEGRADOR"
for p in doc.paragraphs:
    if p.text.strip() == target_text:
        # Crear un párrafo nuevo antes de este
        new_p = docx.oxml.parser.OxmlElement("w:p")
        p._element.addprevious(new_p)
        new_para = docx.text.paragraph.Paragraph(new_p, doc.paragraphs[0]._parent)
        run = new_para.add_run(
            "Revisión v3 (2026-08-24): Se aplicaron las decisiones del grill interno "
            "del equipo sobre los requisitos y las métricas. Cambios principales: "
            "RF-10 pasa a destino favorito, RF-12 añade notificación a usuarios "
            "relevantes, RF-13 lista 5 métricas, RNF-01 sube el tiempo a p95 < 5 s y "
            "añade modo offline básico (solo ruta, sin tiempo de llegada), RNF-02 "
            "baja a 95% mensual, RNF-04 aclara que el login es 1 vez/semana, "
            "RNF-05 se reduce a herramientas en CI, RNF-06 explicita bcrypt, "
            "RNF-07 amplía a Chrome + Safari + escritorio, RNF-08 se elimina, y la "
            "tabla de métricas queda con 5 indicadores (uno por atributo ISO 25010), "
            "eliminando las 3 métricas propias."
        )
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = docx.shared.RGBColor(0x55, 0x55, 0x55)
        break

doc.save(DST)
print(f"v3 saved at: {DST}")
