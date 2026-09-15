"""Corrige RNF-01, RNF-02 y RNF-03 en actividad 3 ing soft_APA v3.docx.

Grill de continuación (2026-09-14) — RF-13, RF-14 y RNF-01..07. Estos tres
cambios salen de contradicciones detectadas entre el .docx y el CLAUDE.md:

- RNF-01: el .docx seguía en "p95 < 5 s" (tabla RNF y tabla de métricas) pese a
  que el CLAUDE.md ya declaraba 8 s como "comprometido en la v3" desde el
  2026-08-26. La memoria del proyecto decía que el .docx se había corregido
  ese día — no es cierto, seguía en 5 s. Se corrige ahora a 8 s en ambas
  tablas.
- RNF-02: asumía un backend corriendo 24/7 en un servidor (UptimeRobot +
  RTO/RPO clásicos). El equipo decidió (2026-09-14) desplegar el frontend
  Angular en Vercel y el backend Spring Boot en Render (tier gratis), que se
  duerme tras 15 min de inactividad con cold start de 30-60 s. RNF-02 se
  reescribe para reflejar esa realidad en vez de un servidor siempre activo.
- RNF-03: no tenía número ni instrumento (violaba la propia regla del curso
  de no enunciar un atributo de calidad sin el número que lo verifica). Se
  alinea con el número que el CLAUDE.md ya usa para mantenibilidad: 0
  violaciones hexagonales en CI + ≥70% cobertura en domain/, vía ArchUnit +
  JaCoCo.
"""

import shutil
from pathlib import Path

import docx
from docx.shared import Pt

SRC = Path(__file__).parent / "actividad 3 ing soft_APA v3.docx"
BACKUP = Path(__file__).parent / "actividad 3 ing soft_APA v3.BACKUP-PRE-RNF-FIX.docx"

if not BACKUP.exists():
    shutil.copyfile(SRC, BACKUP)
    print(f"Backup creado: {BACKUP}")
else:
    print(f"Backup ya existe, no se sobreescribe: {BACKUP}")

doc = docx.Document(SRC)


def rewrite_cell(cell, text):
    for p in list(cell.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run(text)
    run.font.size = Pt(10)


# -------------------------------------------------------------------
# Tabla 2 (RNF) — filas 1 (RNF-01), 2 (RNF-02), 3 (RNF-03).
# -------------------------------------------------------------------
t2 = doc.tables[2]

rewrite_cell(
    t2.rows[1].cells[2],
    "El sistema debe responder las consultas de ruta con un p95 menor a 8 segundos "
    "medido extremo a extremo (cliente → backend → Google Maps → render) en una "
    "red 4G típica, verificado con Spring Boot Actuator + Micrometer. Se cachean "
    "los tiempos de viaje por par origen-destino con un TTL de 5 minutos para no "
    "exceder ese límite. La aplicación ofrece un modo offline básico que muestra "
    "la ruta sugerida y los barrios por los que pasa (calculada por distancia "
    "geométrica sobre el trazado GeoJSON), sin incluir el cálculo de tiempo de "
    "llegada, el cual solo se obtiene en línea vía Google Maps.",
)

rewrite_cell(
    t2.rows[2].cells[2],
    "El backend se despliega en Render (tier gratis) y el frontend en Vercel. El "
    "tier gratis de Render suspende el servicio tras 15 minutos sin tráfico y "
    "tarda entre 30 y 60 segundos en reactivarse ante la siguiente petición "
    "(cold start); esa reactivación no cuenta como caída para efectos de esta "
    "métrica. Medido con UptimeRobot (latido cada 5 minutos) en horario hábil de "
    "lunes a viernes entre 7:00 y 21:00 hora Colombia, el sistema debe responder "
    "correctamente (incluyendo el tiempo de cold start) en al menos un 95% de los "
    "latidos registrados. RPO < 24 horas con backup diario automatizado de "
    "PostgreSQL (Supabase). No se compromete un RTO de servidor siempre activo, "
    "porque el tier gratis no lo sostiene.",
)

rewrite_cell(
    t2.rows[3].cells[2],
    "El backend debe mantener una separación estricta entre las capas del "
    "hexágono (dominio, aplicación, adaptadores), sin que la lógica de negocio "
    "dependa directamente de frameworks externos: 0 violaciones de la "
    "arquitectura hexagonal detectadas por ArchUnit en CI, y al menos 70% de "
    "cobertura de tests unitarios en el paquete domain/, medida con JaCoCo.",
)

# -------------------------------------------------------------------
# Tabla 3 (Métricas) — filas 1 (rendimiento) y 2 (disponibilidad).
# -------------------------------------------------------------------
t3 = doc.tables[3]

rewrite_cell(
    t3.rows[1].cells[1],
    "p95 < 8 s en /api/routes/search (medido extremo a extremo, cliente → "
    "backend → Google Maps → render) + LCP < 2.5 s en Lighthouse mobile con "
    "throttling 'Slow 4G'. Caché de 5 min por par origen-destino.",
)

rewrite_cell(
    t3.rows[2].cells[1],
    "≥ 95% de los latidos de monitoreo responden correctamente en horario hábil "
    "(lun–vie 7:00–21:00), incluyendo el cold start del tier gratis de Render "
    "(hasta 60 s tras 15 min de inactividad); RPO < 24 horas.",
)

rewrite_cell(
    t3.rows[2].cells[2],
    "UptimeRobot (latido cada 5 minutos) contra el backend en Render + backup "
    "diario automatizado de PostgreSQL (Supabase) y restore probado al cierre de "
    "cada sprint.",
)

# -------------------------------------------------------------------
# Nota de revisión.
# -------------------------------------------------------------------
body = doc.element.body
target_text = "DEFINICIÓN DEL PROYECTO INTEGRADOR"
for p in doc.paragraphs:
    if p.text.strip().startswith("Revisión v3 (2026-08-24)"):
        new_p = docx.oxml.parser.OxmlElement("w:p")
        p._element.addnext(new_p)
        new_para = docx.text.paragraph.Paragraph(new_p, doc.paragraphs[0]._parent)
        run = new_para.add_run(
            "Revisión v3.1 (2026-09-14): grill de continuación sobre RF-13, RF-14 "
            "y RNF-01..07 (backlog completo en docs/backlog/). Se corrige RNF-01 a "
            "p95 < 8 s (el .docx seguía en 5 s pese a que el CLAUDE.md ya declaraba "
            "8 s desde el 26-ago; la tabla de métricas tenía el mismo desfase). Se "
            "reescribe RNF-02 sobre el despliegue real decidido esta sesión: "
            "frontend en Vercel, backend en Render (tier gratis, con cold start de "
            "hasta 60 s tras inactividad — no un servidor 24/7 clásico). Se añade "
            "número e instrumento a RNF-03 (0 violaciones hexagonales en CI + "
            "≥70% cobertura en domain/, vía ArchUnit + JaCoCo), que antes no tenía "
            "ninguno de los dos."
        )
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = docx.shared.RGBColor(0x55, 0x55, 0x55)
        break

doc.save(SRC)
print(f"v3.docx corregido y guardado: {SRC}")
